import json, re, collections, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
S = json.load(open('struct2.json', encoding='utf-8'))
CARDS = {c['no']: c for c in json.load(open('cards.json', encoding='utf-8'))}
R = {r['n']: r['t'] for r in json.load(open('rules2.json', encoding='utf-8'))['rules'] if r['l'] >= 0}
QI = json.load(open('qa_items.json', encoding='utf-8'))
qa_n = collections.Counter(q['unit'] for q in QI if q['unit'] is not None)
cn = lambda no: re.sub(r'^\S+\s*', '', CARDS[no]['name']) if no in CARDS else ''
rl = lambda ns: ', '.join(f'{n} {R.get(n, "")[:30]}' for n in ns)

wb = Workbook()
HDR = Font(name='맑은 고딕', size=10, bold=True, color='FFFFFF'); BODY = Font(name='맑은 고딕', size=10)
FILL = PatternFill('solid', fgColor='1F4E78'); TH = Side(style='thin', color='BFBFBF'); BD = Border(left=TH, right=TH, top=TH, bottom=TH)
def sheet(ws, headers, rows, widths):
    ws.append(headers)
    for r in rows: ws.append(r)
    for c, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(c)].width = w
    for i, row in enumerate(ws.iter_rows(), 1):
        for cell in row:
            cell.border = BD
            if i == 1: cell.font, cell.fill = HDR, FILL; cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            else: cell.font = BODY; cell.alignment = Alignment(vertical='top', wrap_text=True)
    ws.freeze_panes = 'B2'; ws.auto_filter.ref = ws.dimensions

# 1. 분류 요약 (15-8)
ws = wb.active; ws.title = '분류 요약'
cls = collections.Counter(s['분류'] for s in S)
desc = {'유발형': '15-8-3 유발 조건(유발 타이밍·「~했을 때」)을 만족하면 유발 → 발휘 대기 → 발휘', '상재형': '15-8-2 유발하지 않고 조건을 만족하는 동안 계속 발휘',
        '기동형': '15-8-4 【메인】 — 턴 플레이어가 메인 페이즈에 선언 (【카운터】 포함)', '즉시형': '15-8-5 「~할 때」·「떠날 때」 — 요인 직전에 끼어들어 발휘',
        '키워드 효과': '16장 키워드 효과(《블로커》 등 단독 표기)', '카드 정보(효과 아님)': '2-3 진화 조건·디지크로스·어셈블리·〈룰〉·〈링크〉 — 효과가 아닌 카드 정보'}
rows = [[k, v, desc.get(k, ''), len({s['group'] for s in S if s['분류'] == k})] for k, v in cls.most_common()]
sheet(ws, ['분류', '효과 수', '룰북 정의', '구조 그룹 수'], rows, [22, 10, 80, 12])

# 2. 구조 그룹
G = collections.defaultdict(list)
for s in S: G[s['group']].append(s)
rows = []
for k, ss in sorted(G.items(), key=lambda kv: -len(kv[1])):
    parts = k.split(' | ')
    ex = min(ss, key=lambda s: len(s['text']))
    rules = collections.Counter(n for s in ss for _, _, n in s['rules'])
    rows.append([k, parts[0], parts[1] if len(parts) > 1 else '', parts[2].replace('처리: ', '') if len(parts) > 2 else '', len(ss), len({s['card'] for s in ss}),
                 sum(qa_n[s['i']] for s in ss), sum(1 for s in ss if s['처리 조건']), sum(1 for s in ss if s['임의 처리 조건']), ex['text'],
                 rl([n for n, c in rules.most_common() if c >= 0.5 * len(ss)][:8]), ', '.join(sorted({s['card'] for s in ss})[:30])])
ws = wb.create_sheet('구조 그룹')
sheet(ws, ['구조 그룹', '분류', '타이밍 / 유발 조건', '주 처리', '효과 수', '카드 수', '연결된 Q&A', '처리 조건 있음', '임의 처리 조건 있음', '대표 문구(가장 짧은 것)', '공통 룰 조항', '카드 목록(앞 30)'],
      rows, [60, 10, 34, 16, 8, 8, 9, 9, 9, 60, 60, 40])

# 3. 효과 구조 (one row per effect)
rows = []
for s in S:
    steps = ' ▶ '.join(f"[{'임의' if p['임의'] else '강제'}{('·' + p['scope']) if p['scope'] else ''}{('·' + p['dur']) if p['dur'] else ''}] {'/'.join(p['acts']) or '?'}" for p in s.get('처리', []))
    rows.append([s['card'], cn(s['card']), s['출처'], s['분류'], s.get('세부', ''), ' '.join('【' + t + '】' for t in s.get('타이밍', [])), ' / '.join(s.get('아이콘', [])),
                 s.get('유발 키', ''), s.get('유발 조건', ''), ' / '.join(s.get('처리 조건', [])), ' / '.join(c['text'] for c in s.get('임의 처리 조건', [])),
                 steps, s['group'], qa_n[s['i']], rl(list(dict.fromkeys(n for _, _, n in s['rules']))), s['text']])
ws = wb.create_sheet('효과 구조')
sheet(ws, ['카드 번호', '카드명', '출처', '분류(15-8)', '세부', '효과 타이밍(15-16)', '아이콘(15-14)', '유발 조건 요약(15-5)', '유발 조건 원문', '처리 조건(15-6)', '임의 처리 조건(15-7)',
           '처리(15-1-2 순서·15-9 강제/임의·15-11 단일/전체)', '구조 그룹', 'Q&A 수', '부품별 룰 조항', '원문'],
      rows, [11, 16, 14, 12, 14, 16, 10, 30, 40, 36, 40, 50, 50, 7, 70, 80])

# 4. 부품 → 조항
cnt = collections.Counter((p, v if p not in ('처리 조건', '임의 처리 조건', '유발 조건') or len(v) < 25 else '(문장)', n) for s in S for p, v, n in s['rules'])
rows = [[p, v, n, R.get(n, ''), c] for (p, v, n), c in sorted(cnt.items(), key=lambda kv: (kv[0][0], -kv[1]))]
ws = wb.create_sheet('부품과 조항')
sheet(ws, ['부품', '값', '조항', '조항 제목', '효과 수'], rows, [14, 34, 10, 40, 9])

ws = wb.create_sheet('안내', 0)
info = [['디지몬 카드 효과 — 룰북 구조 분류 v4', ''], ['만든 날', '2026-10-17'],
        ['기준', '종합 룰 Ver.4.3 15장 「효과의 룰」: 출처(2-3-3·15-3·15-16-10) → 분류 4종(15-8) → 효과 타이밍(15-16)·아이콘(15-14) → 유발 조건(15-5) → 처리 조건(15-6) → 임의 처리 조건(15-7) → 처리(15-1-2 순서, 15-9 강제/임의, 15-10 대상, 15-11 단일/전체, 15-12 정보 추가/변경, 15-13 부여, 18-2 치환)'],
        ['분류 판정 규칙', '【메인】【카운터】=기동형 / 유발 타이밍(【등장 시】 등)·「~했을 때」=유발형 / 「~할 때」「떠날 때」=즉시형 / 【자신의 턴】 등 지속 타이밍만 있고 유발 조건 없음·타이밍 없음=상재형. 「있을 때·이하일 때」 같은 상태는 처리 조건.'],
        ['구조 그룹', '분류 · (유발 조건 또는 타이밍) · 주 처리 가 같은 효과 묶음. 예) 유발형 | 유발 조건: 이 카드 · 진화원에서 · 효과로 · 파기됨 | 처리: 메모리 증가 = EX8-005 고로몬 + BT10-072 바그라군 …'],
        ['효과 수', "=COUNTA('효과 구조'!A:A)-1"], ['구조 그룹 수', "=COUNTA('구조 그룹'!A:A)-1"],
        ['시트', '분류 요약 / 구조 그룹 / 효과 구조(효과마다 부품 분해) / 부품과 조항(부품별로 연결한 룰 조항)'],
        ['검토 부탁', '① 분류 4종 판정이 틀린 효과 ② 유발 조건·처리 조건·임의 처리 조건이 잘못 잘린 효과 ③ 같은 구조 그룹이어야 하는데 갈린(또는 그 반대) 사례']]
for r in info: ws.append(r)
ws.column_dimensions['A'].width = 18; ws.column_dimensions['B'].width = 120
for i, row in enumerate(ws.iter_rows(), 1):
    for c in row: c.font = Font(name='맑은 고딕', size=12 if i == 1 else 10, bold=(i == 1 or c.column == 1)); c.alignment = Alignment(vertical='top', wrap_text=True)
os.makedirs(os.environ.get('OUT_DIR', 'out'), exist_ok=True); path = os.path.join(os.environ.get('OUT_DIR', 'out'), '카드효과_룰북구조_v4.xlsx')
wb.save(path); print(path)
