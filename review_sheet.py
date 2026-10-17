import os
"""Gold-standard review sheet: 300 effects stratified by the automatic class, with O/X dropdowns per component
and a summary sheet that reweights per-class accuracy to the population distribution."""
import json, random, re, collections
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule

S = json.load(open('struct2.json', encoding='utf-8'))
CARDS = {c['no']: c for c in json.load(open('cards.json', encoding='utf-8'))}
cn = lambda no: re.sub(r'^\S+\s*', '', CARDS[no]['name']) if no in CARDS else ''
QUOTA = {'유발형': 110, '상재형': 50, '기동형': 45, '즉시형': 45, '키워드 효과': 25, '카드 정보(효과 아님)': 25}
POP = collections.Counter(s['분류'] for s in S)
random.seed(20261017)
sample = []
for cls, k in QUOTA.items():
    pool = [s for s in S if s['분류'] == cls]
    # spread over sources and timings: sort by (출처, 타이밍) buckets then sample round-robin
    buckets = collections.defaultdict(list)
    for s in pool: buckets[(s['출처'], tuple(s.get('타이밍', []))[:1])].append(s)
    for b in buckets.values(): random.shuffle(b)
    keys = sorted(buckets, key=lambda x: -len(buckets[x])); pick = []
    while len(pick) < k and any(buckets[x] for x in keys):
        for x in keys:
            if buckets[x] and len(pick) < k: pick.append(buckets[x].pop())
    sample += pick
random.shuffle(sample)

def steps(s):
    return '\n'.join(f"{k + 1}) [{'임의' if p['임의'] else '강제'}{('·' + p['scope']) if p['scope'] else ''}{('·' + p['dur']) if p['dur'] else ''}] {' / '.join(p['acts']) or '?'}"
                     for k, p in enumerate(s.get('처리', [])))

wb = Workbook()
ws = wb.active; ws.title = '검수'
H = ['번호', '카드', '카드명', '출처', '원문', '자동: 분류', '① 분류 맞음?', '정답 분류(틀리면)', '자동: 타이밍·아이콘', '자동: 유발 조건', '② 유발 조건 맞음?',
     '자동: 처리 조건', '③ 처리 조건 맞음?', '자동: 임의 처리 조건', '④ 임의 처리 조건 맞음?', '자동: 처리 단계', '⑤ 처리 맞음?', '메모']
ws.append(H)
for k, s in enumerate(sample, 1):
    tm = ' '.join('【' + t + '】' for t in s.get('타이밍', [])) + (' ' + ' '.join('[' + i + ']' for i in s.get('아이콘', [])) if s.get('아이콘') else '')
    trig = (s.get('유발 키') or '') + (('\n원문: ' + s['유발 조건']) if s.get('유발 조건') else '')
    ws.append([k, s['card'], cn(s['card']), s['출처'], s['text'], s['분류'] + (f" ({s['세부']})" if s.get('세부') else ''), '', '', tm.strip(),
               trig or '(없음)', '', '\n'.join(s.get('처리 조건', [])) or '(없음)', '', '\n'.join(c['text'] for c in s.get('임의 처리 조건', [])) or '(없음)', '',
               steps(s) or '(없음)', '', ''])
n = len(sample) + 1
widths = [6, 11, 14, 12, 60, 18, 9, 16, 16, 34, 9, 30, 9, 34, 9, 40, 9, 24]
HDR = Font(name='맑은 고딕', size=10, bold=True, color='FFFFFF'); BODY = Font(name='맑은 고딕', size=10)
FILL = PatternFill('solid', fgColor='1F4E78'); INPUT = PatternFill('solid', fgColor='FFF2CC')
TH = Side(style='thin', color='BFBFBF'); BD = Border(left=TH, right=TH, top=TH, bottom=TH)
for c, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(c)].width = w
INPUT_COLS = [7, 8, 11, 13, 15, 17, 18]
for r, row in enumerate(ws.iter_rows(min_row=1, max_row=n), 1):
    for cell in row:
        cell.border = BD
        if r == 1: cell.font, cell.fill = HDR, FILL; cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        else:
            cell.font = BODY; cell.alignment = Alignment(vertical='top', wrap_text=True)
            if cell.column in INPUT_COLS: cell.fill = INPUT; cell.alignment = Alignment(horizontal='center', vertical='top', wrap_text=True)
ws.freeze_panes = 'F2'; ws.auto_filter.ref = f'A1:R{n}'
ox = DataValidation(type='list', formula1='"O,X,해당없음"', allow_blank=True, showErrorMessage=True, errorTitle='입력', error='O, X, 해당없음 중에서 고르세요')
cls_dv = DataValidation(type='list', formula1='"유발형,상재형,기동형,즉시형,키워드 효과,카드 정보(효과 아님)"', allow_blank=True)
ws.add_data_validation(ox); ws.add_data_validation(cls_dv)
for col in ('G', 'K', 'M', 'O', 'Q'): ox.add(f'{col}2:{col}{n}')
cls_dv.add(f'H2:H{n}')
red = PatternFill('solid', fgColor='F8CBAD'); green = PatternFill('solid', fgColor='C6EFCE')
for col in ('G', 'K', 'M', 'O', 'Q'):
    ws.conditional_formatting.add(f'{col}2:{col}{n}', CellIsRule(operator='equal', formula=['"X"'], fill=red))
    ws.conditional_formatting.add(f'{col}2:{col}{n}', CellIsRule(operator='equal', formula=['"O"'], fill=green))

# ---- 집계 ----
sm = wb.create_sheet('집계')
sm.append(['분류(자동)', '전체 효과 수', '표본 수', '검수 완료(①)', '① 분류 정확도', '② 유발 조건', '③ 처리 조건', '④ 임의 처리 조건', '⑤ 처리', '모집단 가중치'])
R = f'검수!$F$2:$F${n}'
acc = lambda col, c: f'=IFERROR(COUNTIFS({R},"{c}*",검수!${col}$2:${col}${n},"O")/(COUNTIFS({R},"{c}*",검수!${col}$2:${col}${n},"O")+COUNTIFS({R},"{c}*",검수!${col}$2:${col}${n},"X")),"")'
for i, c in enumerate(QUOTA, 2):
    sm.append([c, POP[c], f'=COUNTIF({R},"{c}*")', f'=COUNTIFS({R},"{c}*",검수!$G$2:$G${n},"O")+COUNTIFS({R},"{c}*",검수!$G$2:$G${n},"X")',
               acc('G', c), acc('K', c), acc('M', c), acc('O', c), acc('Q', c), f'=B{i}/SUM($B$2:$B$7)'])
last = len(QUOTA) + 1
sm.append(['표본 전체(가중 안 함)', f'=SUM(B2:B{last})', f'=SUM(C2:C{last})', f'=SUM(D2:D{last})'] +
          [f'=IFERROR(COUNTIF(검수!${c}$2:${c}${n},"O")/(COUNTIF(검수!${c}$2:${c}${n},"O")+COUNTIF(검수!${c}$2:${c}${n},"X")),"")' for c in ('G', 'K', 'M', 'O', 'Q')] + [''])
sm.append(['모집단 추정(분류 분포로 가중)', '', '', ''] +
          [f'=IFERROR(SUMPRODUCT(--({L}2:{L}{last}<>""),{L}2:{L}{last},$J$2:$J${last})/SUMPRODUCT(--({L}2:{L}{last}<>""),$J$2:$J${last}),"")' for L in ('E', 'F', 'G', 'H', 'I')] + [''])
sm.append([])
sm.append(['혼동표(행=자동 분류, 열=정답 분류; 분류가 X인 행만)'])
labels = list(QUOTA)
sm.append([''] + labels)
start = sm.max_row + 1
for c in labels:
    sm.append([c] + [f'=COUNTIFS({R},"{c}*",검수!$G$2:$G${n},"X",검수!$H$2:$H${n},"{d}")' for d in labels])
for row in sm.iter_rows():
    for cell in row:
        cell.font = Font(name='맑은 고딕', size=10, bold=cell.row == 1 or cell.column == 1); cell.border = BD
        cell.alignment = Alignment(vertical='top', wrap_text=True)
for r in range(2, last + 3):
    for c in range(5, 11): sm.cell(r, c).number_format = '0.0%'
for c, w in enumerate([30, 12, 10, 12, 14, 12, 12, 14, 10, 12], 1): sm.column_dimensions[get_column_letter(c)].width = w

# ---- 안내 ----
gd = wb.create_sheet('안내', 0)
for r in [['룰북 구조 분류 — 정답셋 검수 (표본 300)'],
          ['목적', '자동 분해(종합 룰 15장 기준)가 맞는지 사람이 판정한 정답셋을 만들어, 논문에서 분류·부품 추출의 정확도를 보고합니다.'],
          ['표본', '자동 분류별 층화 추출: 유발형 110, 상재형 50, 기동형 45, 즉시형 45, 키워드 효과 25, 카드 정보 25 (seed 20261017). 드문 분류도 충분히 검증되도록 많이 뽑았고, 전체 정확도는 「집계」 시트에서 실제 분포로 다시 가중합니다.'],
          ['하는 일', '노란 칸만 채우면 됩니다. ①~⑤ 각 항목이 맞으면 O, 틀리면 X, 그 효과에 해당하지 않으면 「해당없음」. ① 분류가 X면 「정답 분류」에서 맞는 분류를 고르세요. 이상한 점은 메모에.'],
          ['판정 기준', '① 분류(15-8): 【메인】【카운터】=기동형 / 유발 타이밍·「~했을 때」=유발형 / 「~할 때」「떠날 때」=즉시형 / 지속 타이밍만·타이밍 없음=상재형 / 진화 조건·디지크로스·〈룰〉=카드 정보\n② 유발 조건(15-5): 효과를 유발하는 조건을 맞게 잡았는가(요약·원문)\n③ 처리 조건(15-6): 「~라면」「~인 동안」을 빠짐없이·과하지 않게 잡았는가\n④ 임의 처리 조건(15-7): 「~하는 것으로」 앞 부분 전체를 잡았는가\n⑤ 처리: 단계 나눔·처리 종류·강제/임의·대상이 대체로 맞는가'],
          ['해당없음', '그 효과에 없는 부품을 자동 결과도 「(없음)」으로 두었다면 O입니다. 있어야 하는데 「(없음)」이면 X입니다. 키워드 효과·카드 정보는 ①만 판정하고 ②~⑤는 「해당없음」.'],
          ['진행', '일부만 해도 집계가 자동으로 계산됩니다. 100개씩 나눠 하셔도 됩니다.']]:
    gd.append(r)
gd.column_dimensions['A'].width = 14; gd.column_dimensions['B'].width = 110
for row in gd.iter_rows():
    for cell in row: cell.font = Font(name='맑은 고딕', size=12 if cell.row == 1 else 10, bold=cell.row == 1 or cell.column == 1); cell.alignment = Alignment(vertical='top', wrap_text=True)
out = os.path.join('paper', '정답셋_검수_300.xlsx')
wb.save(out)
json.dump([s['i'] for s in sample], open(os.path.join('paper', 'results', 'gold_sample_ids.json'), 'w'))
print(out, collections.Counter(s['분류'] for s in sample))
